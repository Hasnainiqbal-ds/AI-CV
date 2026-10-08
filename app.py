"""ATS Resume Checker - Streamlit + Google Gemini Flash.

Upload a resume (PDF, DOCX or TXT) and get an ATS score with
section-wise feedback and concrete improvements.
"""

import io
import json
import os
import re
import time

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# ----------------------------- Configuration ----------------------------- #

DEFAULT_MODEL = os.environ.get("GEMINI_MODEL", "gemini-3.7-flash")
MODEL_CHOICES = [DEFAULT_MODEL, "gemini-3.6-flash", "gemini-3.5-flash"]
MODEL_CHOICES = list(dict.fromkeys(MODEL_CHOICES))  # remove duplicates, keep order

RETRIES_PER_MODEL = 3  # attempts per model before trying the next one
BACKOFF_SECONDS = 2  # wait 2s, then 4s, between attempts

MAX_FILE_MB = 5
MAX_CHARS = 20000  # resume text sent to the model
MIN_CHARS = 150  # below this the file is probably scanned / empty

# Weights used to compute the final ATS score (sum = 100).
WEIGHTS = {
    "keyword_match": 30,
    "formatting": 20,
    "sections": 15,
    "impact": 20,
    "readability": 15,
}

LABELS = {
    "keyword_match": "Keywords & relevance",
    "formatting": "ATS-friendly formatting",
    "sections": "Standard sections",
    "impact": "Impact & achievements",
    "readability": "Readability & language",
}


# --------------------------- File text extraction ------------------------- #

def extract_text(data: bytes, filename: str) -> str:
    """Return plain text from a PDF, DOCX or TXT file."""
    name = filename.lower()

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password protected.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n".join(pages).strip()

    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip()

    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()

    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


# ------------------------------ Prompt + AI ------------------------------- #

def build_prompt(resume_text: str, job_description: str) -> str:
    jd_block = (
        f"JOB DESCRIPTION:\n{job_description.strip()}\n"
        if job_description.strip()
        else "JOB DESCRIPTION: (not provided - judge keywords for the general "
        "role this resume targets)\n"
    )
    return f"""You are an expert ATS (Applicant Tracking System) analyst and resume coach.
Analyse the resume below and return ONLY a valid JSON object, no markdown, no extra text.

Score each category from 0 to 100 (integers):
- keyword_match: relevant skills/keywords, matched to the job description if given
- formatting: ATS-friendly structure (clear headings, no tables/columns/graphics issues, consistent dates, contact info)
- sections: presence of standard sections (contact, summary, experience, education, skills, projects/certifications)
- impact: use of action verbs and measurable achievements instead of duties
- readability: clarity, grammar, concise bullets, sensible length

Be honest and strict. Do not inflate scores.

JSON schema:
{{
  "scores": {{"keyword_match": 0, "formatting": 0, "sections": 0, "impact": 0, "readability": 0}},
  "summary": "2-3 sentence overall assessment",
  "strengths": ["..."],
  "weaknesses": ["..."],
  "missing_keywords": ["..."],
  "improvements": [
    {{"priority": "High|Medium|Low", "section": "...", "issue": "...", "fix": "..."}}
  ],
  "rewrite_examples": [
    {{"before": "weak bullet from the resume", "after": "stronger rewritten bullet"}}
  ]
}}

Give 3-6 strengths, 3-6 weaknesses, up to 12 missing keywords, 5-10 improvements and 2-4 rewrite examples.
Only use facts that exist in the resume. Never invent numbers in rewrite examples; use placeholders like [X%] where a metric is needed.

{jd_block}
RESUME:
{resume_text[:MAX_CHARS]}
"""


def parse_json(raw: str) -> dict:
    """Parse JSON from the model, tolerating ``` fences or stray text."""
    text = (raw or "").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.IGNORECASE).strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        start, end = text.find("{"), text.rfind("}")
        if start != -1 and end > start:
            return json.loads(text[start : end + 1])
        raise ValueError("The AI response was not valid JSON. Please try again.")


def _clamp(value, low=0, high=100) -> int:
    try:
        return max(low, min(high, int(round(float(value)))))
    except (TypeError, ValueError):
        return 0


def _str_list(value) -> list:
    if not isinstance(value, list):
        return []
    return [str(v).strip() for v in value if str(v).strip()]


def normalize_result(data: dict) -> dict:
    """Validate the model output and compute the final weighted ATS score."""
    if not isinstance(data, dict):
        raise ValueError("Unexpected AI response format.")

    raw_scores = data.get("scores") if isinstance(data.get("scores"), dict) else {}
    scores = {key: _clamp(raw_scores.get(key, 0)) for key in WEIGHTS}
    overall = _clamp(sum(scores[k] * w for k, w in WEIGHTS.items()) / 100)

    improvements = []
    for item in data.get("improvements", []) if isinstance(data.get("improvements"), list) else []:
        if not isinstance(item, dict):
            continue
        priority = str(item.get("priority", "Medium")).strip().capitalize()
        if priority not in ("High", "Medium", "Low"):
            priority = "Medium"
        improvements.append(
            {
                "priority": priority,
                "section": str(item.get("section", "General")).strip() or "General",
                "issue": str(item.get("issue", "")).strip(),
                "fix": str(item.get("fix", "")).strip(),
            }
        )
    order = {"High": 0, "Medium": 1, "Low": 2}
    improvements.sort(key=lambda i: order[i["priority"]])

    rewrites = []
    for item in data.get("rewrite_examples", []) if isinstance(data.get("rewrite_examples"), list) else []:
        if isinstance(item, dict) and item.get("before") and item.get("after"):
            rewrites.append({"before": str(item["before"]).strip(), "after": str(item["after"]).strip()})

    return {
        "overall": overall,
        "scores": scores,
        "summary": str(data.get("summary", "")).strip(),
        "strengths": _str_list(data.get("strengths")),
        "weaknesses": _str_list(data.get("weaknesses")),
        "missing_keywords": _str_list(data.get("missing_keywords")),
        "improvements": improvements,
        "rewrites": rewrites,
    }


def _error_code(exc: Exception):
    """Best-effort HTTP status code from an API error."""
    code = getattr(exc, "code", None)
    if isinstance(code, int):
        return code
    match = re.search(r"\b(4\d\d|5\d\d)\b", str(exc))
    return int(match.group(1)) if match else None


def is_transient(exc: Exception) -> bool:
    """Temporary server problems: overload (503), rate limit (429), timeouts."""
    text = str(exc).upper()
    return _error_code(exc) in (429, 500, 502, 503, 504) or any(
        w in text for w in ("UNAVAILABLE", "OVERLOADED", "HIGH DEMAND", "RESOURCE_EXHAUSTED", "TIMED OUT")
    )


def is_model_missing(exc: Exception) -> bool:
    return _error_code(exc) == 404 or "NOT FOUND" in str(exc).upper()


def analyze_resume(api_key: str, model: str, resume_text: str, job_description: str) -> dict:
    """Call Gemini with retries; fall back to other Flash models if one is busy."""
    client = genai.Client(api_key=api_key)
    prompt = build_prompt(resume_text, job_description)
    config = types.GenerateContentConfig(temperature=0.2, response_mime_type="application/json")

    models_to_try = [model] + [m for m in MODEL_CHOICES if m != model]
    last_error = None

    for current in models_to_try:
        for attempt in range(RETRIES_PER_MODEL):
            try:
                response = client.models.generate_content(model=current, contents=prompt, config=config)
                result = normalize_result(parse_json(response.text))
                result["model_used"] = current
                return result
            except Exception as exc:
                last_error = exc
                if is_model_missing(exc):
                    break  # this model name is not valid, try the next one
                if not is_transient(exc):
                    raise  # bad API key, bad request, invalid JSON, etc.
                if attempt < RETRIES_PER_MODEL - 1:
                    time.sleep(BACKOFF_SECONDS * (2 ** attempt))
                # otherwise loop ends and we move on to the next model

    if last_error is not None and not is_transient(last_error):
        raise last_error
    raise RuntimeError(
        "Gemini is overloaded right now (tried "
        + ", ".join(models_to_try)
        + "). Please wait a minute and try again. Last error: "
        + str(last_error)
    )


# --------------------------------- UI ------------------------------------- #

def get_api_key() -> str:
    """Look for the key in Streamlit secrets, then environment variables."""
    try:
        key = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:  # no secrets file present
        key = ""
    return key or os.environ.get("GEMINI_API_KEY", "")


def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent - likely to pass most ATS filters"
    if score >= 65:
        return "Good - a few fixes will make it stronger"
    if score >= 50:
        return "Average - needs noticeable improvement"
    return "Weak - likely to be filtered out; needs major work"


def render_results(result: dict) -> None:
    st.subheader("Your ATS Score")
    if result.get("model_used"):
        st.caption(f"Analyzed with {result['model_used']}")
    left, right = st.columns([1, 2])
    with left:
        st.metric("Overall ATS score", f"{result['overall']} / 100")
    with right:
        st.progress(result["overall"] / 100)
        st.caption(score_label(result["overall"]))

    if result["summary"]:
        st.info(result["summary"])

    st.subheader("Score breakdown")
    for key, label in LABELS.items():
        value = result["scores"][key]
        st.write(f"**{label}** ({WEIGHTS[key]}% weight): {value}/100")
        st.progress(value / 100)

    col1, col2 = st.columns(2)
    with col1:
        st.subheader("Strengths")
        for item in result["strengths"] or ["No strengths reported."]:
            st.write(f"- {item}")
    with col2:
        st.subheader("Weaknesses")
        for item in result["weaknesses"] or ["No weaknesses reported."]:
            st.write(f"- {item}")

    if result["missing_keywords"]:
        st.subheader("Missing keywords")
        st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))

    st.subheader("Recommended improvements")
    icons = {"High": "🔴", "Medium": "🟠", "Low": "🟢"}
    if not result["improvements"]:
        st.write("No improvements reported.")
    for item in result["improvements"]:
        title = f"{icons[item['priority']]} {item['priority']} - {item['section']}"
        with st.expander(title):
            st.markdown(f"**Issue:** {item['issue']}")
            st.markdown(f"**Fix:** {item['fix']}")

    if result["rewrites"]:
        st.subheader("Example rewrites")
        for item in result["rewrites"]:
            st.markdown(f"**Before:** {item['before']}")
            st.markdown(f"**After:** {item['after']}")
            st.divider()

    report = json.dumps(result, indent=2)
    st.download_button(
        "Download report (JSON)",
        data=report,
        file_name="ats_report.json",
        mime="application/json",
    )


def main() -> None:
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Checker")
    st.write("Upload your resume to get an ATS score and clear steps to improve it.")

    with st.sidebar:
        st.header("Settings")
        api_key = get_api_key()
        if api_key:
            st.success("Gemini API key loaded.")
        else:
            api_key = st.text_input(
                "Gemini API key",
                type="password",
                help="Get a free key at https://aistudio.google.com/apikey",
            )
        model = st.selectbox("Gemini model", MODEL_CHOICES)
        st.caption("If a model name is rejected, pick another one.")

    uploaded = st.file_uploader("Upload resume (PDF, DOCX or TXT)", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional, improves keyword matching)",
        height=150,
        placeholder="Paste the job description here...",
    )

    if st.button("Analyze resume", type="primary"):
        if not api_key:
            st.error("Please add your Gemini API key in the sidebar.")
            return
        if uploaded is None:
            st.error("Please upload a resume first.")
            return
        if uploaded.size > MAX_FILE_MB * 1024 * 1024:
            st.error(f"File is too large. Maximum size is {MAX_FILE_MB} MB.")
            return

        try:
            with st.spinner("Reading your resume..."):
                text = extract_text(uploaded.getvalue(), uploaded.name)
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            return

        if len(text) < MIN_CHARS:
            st.error(
                "Very little text was found. If your resume is a scanned image, "
                "export it as a text-based PDF or DOCX and try again."
            )
            return

        try:
            with st.spinner("Analyzing with Gemini..."):
                result = analyze_resume(api_key, model, text, job_description)
        except Exception as exc:
            st.error(f"Analysis failed: {exc}")
            return

        st.session_state["result"] = result

    # Render from session_state so results survive reruns (e.g. clicking Download).
    if "result" in st.session_state:
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
